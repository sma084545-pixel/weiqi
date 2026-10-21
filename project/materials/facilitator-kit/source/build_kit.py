#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build 15 editable A4 DOCX/Markdown documents from one explicit page model."""
from pathlib import Path
import json, zipfile, re, os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET

K=Path(__file__).resolve().parents[1];R=K.parents[1];OUT=K/'output'
D=json.loads((K/'source/kit.json').read_text());catalog=json.loads((R/'docs/merge-c/RESOURCE_CATALOG.json').read_text())
cards={x['id']:x for x in catalog['cards']}
P=lambda t:dict(type='paragraph',text=t)
H=lambda t:dict(type='heading',text=t)
T=lambda heads,rows,widths=None:dict(type='table',headers=heads,rows=rows,widths=widths)
L=lambda texts:dict(type='list',items=texts)
def page(title,*blocks):return dict(title=title,blocks=list(blocks))
def ident(id):return dict(schemaVersion=1,id=id,revision=1,rulesetId='GO-CLUB-AREA-1')

resources=[
 ['B5','完整 5 路示意盘','materials/facilitator-kit/output/T23_board_5x5_A4.pdf','1'],
 ['B9','9 路主盘','materials/pair-boards/output/boards/go-board_9x9-a4_coordinates.pdf','1'],
 ['B13','13 路首选 A3','materials/pair-boards/output/boards/go-board_13x13-a3_coordinates.pdf','1'],
 ['B13T','13 路 A4 拼接','materials/pair-boards/output/boards/go-board_13x13-a3_coordinates_tiled-A4.pdf','1–2'],
 ['S18','18 mm 棋子 每色 90','materials/pair-boards/output/portable-kit/paper-stones_18mm_black-white_180.pdf','1–2'],
 ['O9','可选气覆片','materials/pair-boards/output/overlays/liberty-overlay_9x9-a4.pdf','1'],
 ['RB','基础规则','materials/rule-cards/pdf/main-rule-A5-duplex.pdf','1–2'],
 ['RD','规则详情','materials/rule-cards/pdf/detailed-A5-duplex.pdf','1–2'],
 ['RS','面积计分','materials/rule-cards/pdf/05-score.pdf','1'],
 ['RM','吃子小游戏','materials/rule-cards/pdf/capture-mini-A6-duplex.pdf','1–2'],
 ['CF','挑战卡正面','materials/challenge-cards/output/T21_challenge_cards_front_A4.pdf','1–15'],
 ['CB','挑战卡背面','materials/challenge-cards/output/T21_challenge_cards_back_A4.pdf','1–15'],
 ['CA','教师答案册','materials/challenge-cards/output/T21_teacher_answer_book.pdf','1–33'],
 ['DT','双面套准试印','materials/challenge-cards/output/T21_duplex_test_page1.pdf','1–2'],
 ['DB','可选演示盘拼接','materials/demo-board/board-9x9-a4-tiled.pdf','1–6'],
 ['DG','可选演示盘装配','materials/demo-board/assembly-guide.pdf','1–8'],
]
module_names=[('气实验室','liberty-lab','#/learn/liberties'),('提子实验室','capture-lab','#/learn/captures'),('吃子小游戏','capture-game','#/play/partner'),('完整对弈','full-game','#/play/full-game'),('互动题','puzzle-runner','#/challenges'),('眼专题','life-lab','#/learn/life'),('合作挑战','cooperative-challenge','#/challenges/together'),('围地练习','territory-trainer','#/learn/territory'),('新草稿','teacher-editor','#/teacher/editor'),('文件交接','file-exchange','#/teacher/file-exchange'),('离线与备份','offline-data','#/teacher/offline-data'),('自愿反馈','member-feedback','#/feedback'),('反馈与修订','evidence-log','#/teacher/evidence')]

def cards_line(s):
 return '；'.join(f"{id}（{cards[id]['frontPdfPage']}→{cards[id]['teacherPage']}）" for id in s['puzzles'])
def lesson_line(s):
 if not s['lessons']:return '无必看新微课；可按需回看 T22_All_25_r1.pptx。本次以创作、交流和对弈为主。'
 if s['id']=='T23-ACT-01':return '微课：T22_MicroLesson_01_r1.pptx 第 1–5 页（合并 1–5）；T22_MicroLesson_02_r1.pptx 第 1–2 页（合并 6–7）。第 2 套 3–5 页桌边按需看。'
 return '微课：'+ '；'.join(f"T22_MicroLesson_0{l['number']}_r1.pptx 第 {min(l['pages'])}–{max(l['pages'])} 页（合并第 {(l['number']-1)*5+min(l['pages'])}–{(l['number']-1)*5+max(l['pages'])} 页）" for l in s['lessons'])
def session_page(s):
 g='G＝⌈N/2⌉。每组 B9/B5 各 1 盘、黑白各 90 子、1 份规则卡、2 支笔，所选题卡各 1 套。'
 if s['id'] in ['T23-ACT-02','T23-ACT-05','T23-ACT-06']:g+=' 13 路可选：每组加 B13/B13T、各色 180 子；无大盘继续 9 路。'
 tail=s.get('fixtures',[])
 return page(s['title'],P(s['goal']),P('器材与布置：'+g+' '+s['layout']),P(lesson_line(s)),P('题卡均 r1，括号为正/背页→答案页：'+cards_line(s)+(('；终局 '+ '、'.join(tail)+' r1，见围地练习。') if tail else '。')),P('开场：'+s['opening']),T(['时段','主持与成员动作'],[[f"{sum(x['minutes'] for x in s['timeline'][:i])}–{sum(x['minutes'] for x in s['timeline'][:i+1])} 分",x['title']+'：'+x['action']] for i,x in enumerate(s['timeline'])],[22,152]),L(s['paths']),P('无屏：'+s['noscreen']),P('结束问：'+s['ending']),P('收尾：自愿三项反馈或跳过；活动编号/日期按实际填写，之后清点收纳。'),P('注意：'+s['safeguard']))

prep=page('人数器材和空间准备',P('默认中学生、每次 45 分钟、两人共用棋盘，均为可调整建议。N 为实际到场人数，G＝⌈N/2⌉，备用套数 U 由组织者按预算填写；准备 G＋U 套。奇数人数可一桌三人，轮换下棋、回应与观察，允许先旁观。'),
 T(['物品','数量与低成本替代'],[
 ['基础棋具','G＋U 张 B9 与 G＋U 张 B5；18 mm 黑白子每色 90(G＋U) 枚。S18 每套 2 页 A4，100% 打印。'],
 ['13 路可选','H 个扩展盘用 H 张 B13（A3）或 2H 张 B13T（A4 拼接）；18 mm 棋子每色 180H 枚，S18 每组印 2 套共 4 页。'],
 ['圈气与目标','每组 1 个透明袋/覆片可选；无覆片就用纸环或指尖。空点气环与棋上目标三角分开。'],
 ['纸卡与规则','每组 1 套所选题卡、规则卡。所选题卡占 q 个不同正面页：双面 Gq 张 A4，单面 2Gq 张；背面同页配对。'],
 ['主持参考','1 份手册；答案册可只印对应页。N 张自愿反馈空表；观察表 1 张，实际棋谱每局按手数续页。'],
 ['笔与空间','2G 支笔或可轮换，G 个棋子收纳处；G 个桌面位置。若设 S 个任务站，额外桌面不足就原桌换卡。']],[32,142]),
 P('建议按“接待—双人桌—参考材料—回收”布置。熟练成员分散到桌边，不围住新手代下。显示文字与坐标；允许口述、指点或用较大的 A3 棋盘，不依赖颜色和速度。'),
 P('盘上容量不等于整局手数。提子可回收到供应袋再用，统计另记；面积分不把提子再加一次。纸盘格距与棋子直径须匹配：B9 格距 20 mm 配 18 mm 子，B13 同样；不强行把大棋子挤在小盘上。'),
 P('打印为 A4/A3 原尺寸 100%，先试校准与双面套准页；不使用“适合页面”缩放。纸棋子可以裁切，透明袋可替专用覆膜。T18 演示盘设计可选，不要求购买磁吸材料或特定品牌。'))

rules=page('主持人共同规则速查',P('本项目使用 GO-CLUB-AREA-1 自定义教学规则，不冒称完整中国、日本或 AGA 竞赛规则。活动与题目均可跳过；死活不是开始下棋的先决门槛。'),L([
 '交叉点（intersection）：完整棋盘支持 5/9/13/19 路，入门对弈默认 9 路。5 路要明确标示为完整小盘练习；裁图边缘不是实际盘边。',
 '气（liberty）与连接（connection）：同色上下左右相邻为连通块，斜邻不算；整块邻接空点去重，每个空点只算一次。只剩一气称打吃（atari）。',
 '提子（capture）的顺序：核对轮次、点位与占用，暂放己子；同时移走相邻无气对方块；再查己方整块有气；最后查重复。先提后有气可以合法，真正自杀（suicide）禁止。无效动作不换手，不当作停手。',
 '位置超级劫（positional superko）：落子完成提子后，整个布局不能等于本局路径中任一先前布局，包括起点；只比棋子，不按下一手颜色区别。先用简单劫解释，较长重复同样禁止。',
 '停一手（pass）不受重复检查、不清历史。完整对弈两次连续停手进入终局（endgame）确认，不自动删死子；认输（resign）另形成结果，不编造计分。',
 '双方确认死子与中立点后才最终计分。修改注释取消旧确认。有争议恢复原盘，保留历史/轮次，连续停手数归零；预览移除不是真提子。双方同意撤回时恢复整套状态，程序回看历史局面只读。',
 '面积计分（area scoring）：确认死子移除后，黑活子＋黑围空；白活子＋白围空＋开局贴目（komi）。提子不额外加分；混合边界空区、完全空盘均中立。默认贴目 0、允许和棋，是教学选择，开局前可改。未完成盘面的染色只能作预览。',
 '吃子小游戏（Capture Go）：首个成功提子者立即赢本轮，两停无胜者；不用面积决定小游戏胜负。摆棋（setup）可任意摆/擦，不是合法棋谱；进入试下前检查每块有气、轮次及坐标，从当前盘建立新历史，过去未知，不伪造。',
 '眼（eye）：两个可靠保护的独立真眼是基础活形，两个空点不等于两眼；没有两眼不自动死，双活（seki）存在。未收录合法变化保留未知，不作自动死活判定。'
 ]),P('显示坐标 A–T 跳过 I，行从下往上。9 路左上 A9、右下 J1；底层 row/col 从左上 (0,0) 开始。纸面复核用完整棋盘与逐手记录，不能用落点相同代替整盘重复判断。'))

digital=page('程序准备离线和反馈操作',P('程序是可选工具；核心流程都有实体替代。先按根 README 通过本地服务启动，或使用已配置 HTTPS 地址；没有学校地址时，不编造网页链接。地址末尾的 #/ 路由见资源索引。'),
 H('活动前'),L(['打开“离线与本机数据”，点击“准备或重新核对离线资源”。只有显示“已可离线”才当作准备成功；首次资源未缓存时断网不能承诺可用。',
 '在同一浏览器做一次断网刷新演练，再打开气实验室、题目与保存局面；不成功就先用纸面方案。Service Worker 需要 localhost 或 HTTPS，双击 file:// 没有同等能力。',
 '应用缓存含已安装学习模块和正式题库；本手册、T21 PDF、T22 PPTX 要另存本机。程序使用系统字体；PPTX 打开前按原材料说明安装中文字体或换字体后复核。']),
 H('活动中与收尾'),L(['同屏操作先说当前执色。方向键移动、Enter/Space 选点；手机用轻触，不靠悬停。教师视图可切投影模式；先预测，再显示提示。',
 '收到新版提示时先保存当前对局/草稿，再应用更新，不在未保存编辑中刷新。保存受限或配额提示时先导出完整备份，或用纸面记录，避免清库试错。',
 '“导出完整备份”保存本机记录。恢复先“选择备份文件”并看预览；确认替换会替换本机记录，建议先备份现状。清进度不会删草稿或正式题库；选择性清除要核对范围。',
 '“自愿反馈”先选择/建立实际活动编号与日期。提交后表单清空，点“下一位成员（空表）”；可以跳过。共享设备答复数不等于独立成员数。',
 '“反馈与资源迭代”看实际答复、缺失项和分母。导出前预览删减自由文本；完整备份可能含未删减原文，只作本机恢复，不默认分享。']),
 P('新手课堂只讲操作文字，不讲引擎/schema。0 条记录就保持空态；不载入开发测试答复充当成员反馈。断网或无设备时直接转同号纸卡、纸棋谱与自愿空表。'))

intro=page('围棋社活动主持手册',P('这套手册把六次可调整的 45 分钟活动与现有纸卡、微课和程序连接起来。组织者按实际人数准备，每位成员可以选任务、直接对弈、观战或跳过反馈。用途是提供放松、社交与主动思考的休闲选择。'),P('默认面向中学生，两人一盘；人数、设备、预算与学校场地尚未确认。预计时长未经课堂排练，不保证适合所有班级。六次活动安排与课程六阶段有关，但活动编号不是通关顺序。前三次不要求掌握全部眼与死活。'),
 T(['本册页码','使用内容'],[['2–4','人数器材、共同规则、离线和反馈准备'],['5–10','六次活动，每页均含三种参与路径与无屏操作'],['11–15','四种玩法、常见情况、替代表、同伴帮助与真实记录'],['16–19','实际文件、卡片答案页码与程序入口'],['20','来源、使用边界与交接']],[24,150]),
 P('独立文件：活动 01–06 各一页；替代表、同伴卡、准备收尾、反馈、观察、修订、新成员接待、纸棋谱共八份。全部有 DOCX、Markdown 和 A4 PDF。题卡编号与答案册保持原版本；作者参考先藏起来，成员预测后再按需看。'),
 P('主持前先看准备清单与所选活动页。纸卡布局严格照同号原卡，尺寸、目标、先手与条件不能省；微课仅用于短讲解，实际下棋和交流占主要时间。资源索引的 CF/CB/CA 等短码在本册第 16–17 页定义。'),
 P('交付版本 0.17.0-plus-t23.1，基于 MERGE_C 0.17.0-merge-c.1；程序版本仍为 0.17.0。正式题库 30 题与 6 终局保持冻结 revision 1。本册是设计与制作成果，不是已经开展的活动记录或学校 CAS 认可证明。'))

peer_page=page('同伴帮助和新成员接待',H('帮助者把选择留给伙伴'),L(D['peer']),H('随到随加入'),P('问“想先看一手、找气还是直接下？”不用姓名与班级，找到愿意同行的伙伴即可。给一张规则卡、一个完整棋盘与可选择入口，介绍坐标、轮次和暂停的方法。'),P('最短接待示范：把一颗子放交叉点，沿线指出空点叫气；下一手先由新成员自己选。需要补充提子时只演示一手。等到实际出现劫、结束争议或眼形，再给对应小提示，不把全部死活作为加入门槛。'),P('若新成员加入正在进行的局，不擅自改盘或接管执色；先双方同意观察/接替，或者另开一局。想离开可直接结束活动，保留未完成棋局，不处罚、不要求连续打卡。'))

records=page('真实反馈观察和资源修订',P('活动编号、日期与使用资源版本按实际填写。反馈自愿、无需身份资料；三项是难度感受、最喜欢的活动、希望下次改变什么。选项与程序一致，允许全留空。纸表默认不选导出同意，导出前再查看原文。'),
 P('汇总按实际答复数作分母，单列缺失，不能把提交数当人数。原文可以删减以免分享个人信息，未同意导出的意见不默认进入报告。没有答复就写 0 条；不推断智力、压力治疗、学业或全校效果。'),
 H('教师只写实际看见或听见的事'),P('观察表记录：是否提出理由、是否考虑回应、是否修订、是否让同伴解释，以及确切例子与资源 ID/revision。未观察则留空或标未观察；不从表单填写自动宣布观点优秀或合作真实发生。'),
 H('修订与 CAS 过程'),P('修订表写问题、证据引用、旧 ID/revision、拟改内容、修改日期、验证方法、新版本、实施日期、实施后观察和下一步。正式冻结题保持不变；成员在纸稿或教师编辑器中建立独立本机草稿，保留前后版本。正式内容若将来改动，须新 revision、重新核查并同步纸卡/课件与冻结清单。'),
 P('过程证据分别注明本人设计/修改、同伴贡献、AI 协助、参考资料与真实实施。实际投入时间由本人填写，程序生成耗时不自动换算 CAS 工时。反思可问：实际遇见哪处困难？依据是什么？改变了什么？还有什么未验证？没有实施后反馈就保留未观察，不代写第一人称经历。'),
 P('本包空表只提供结构，没有真实成员评价、名单、照片、出勤或工时。纸面意见可按实际录入反馈界面，组织者不能替成员勾选同意。'))

print_rows=[]
for code,label,path,pages in resources:
 print_rows.append([code,label+'\n'+Path(path).name,pages])
paper=page('纸面资源索引',P('以下均为集成工程可用的实际成品。按 materials/ 子目录查找；完整可点击相对路径在配套 Markdown 与 source/resource-map.json 中。表中页码从 1 开始，单位为 PDF 页。'),T(['短码','文件与用途','页码'],print_rows,[14,143,17]),P('B5 在本包 output/；其余 B/S/O 在 pair-boards/output/；RB/RD/RS/RM 在 rule-cards/pdf/；CF/CB/CA/DT 在 challenge-cards/output/；DB/DG 在 demo-board/。同号卡正背面页相同，每页两题，上下位置以 CARD_INDEX.json 为准。先 100% 试印 DT，再裁切题卡；PDF 不替代可编辑源。'))

digital_index=page('微课题号和程序入口索引',T(['微课','独立文件页与合并页'],[[f"第 {n} 套",f"T22_MicroLesson_0{n}_r1.pptx 第 1–5 页\nT22_All_25_r1.pptx 第 {(n-1)*5+1}–{n*5} 页"] for n in range(1,6)],[24,150]),P('微课在 materials/micro-lessons/output/，同名 PDF 可替换 PPTX。讲稿在 source/speaker-notes.md；教师摘要 T22_Teacher_Summary_5pages_r1.pdf 每页对应一套。本册活动页写明选用页，未引用的页不强制观看。'),T(['程序名称','模块 ID 与路由'],[[name,mid+'  '+route] for name,mid,route in module_names],[32,142]),P('在已启动程序中选择对应文字入口，或替换地址的 #/ 部分。没有部署地址，不提供假网页。六阶段入口 #/route 可自由跳转；题目列表按 ID 选题，没有另造单题深链。END-002/006 在围地练习内切换。所有题目 ID/revision 与卡片答案页在各活动页列明。'))

source_page=page('来源与后续使用',P('六次活动、分流、处理话术、材料配置和空表是本项目原创表达。规则以 docs/RULES.md 的 GO-CLUB-AREA-1 为唯一约定；题目、棋形与微课引用既有 T09/T20/T21/T22 资源，本次没有另造同号题或修改答案。'),
 *[P(x['title']+'\n'+x['url']+'\n访问 '+x['accessed']+'。'+x['use']+x['excluded']) for x in D['sources']],
 P('资料中对成年人/儿童或不同赛规的建议不等于本校情况。本项目不沿用排名门槛、固定贴目或停手交子，也不将协会教学建议当作本社团已经取得的成效。'),
 H('使用前再做现场检查'),P('已完成文件生成、时间核算、资源引用与版面检查。字体与打印边界须在实际设备复核；磁吸、耐用性、双面套准和实际课堂时长仍待实体/课堂试验。没有人类教师审查或真实参与数据，不能把内部核查写成这些证据。'),
 P('下一任务 T24 使用本完整集成归档、handoff/T23.md 与本包 validation/。程序、共同 schema、规则与正式题库保持不变；不用旧材料支线程序替换主线。T24 尚未执行，本手册不代替学校审批或保证 CAS 获认可。'))

checklist=page('活动准备与收尾清单',P('活动编号 __________  日期 __________  主持角色 __________'),H('准备'),L(['[ ] 确认实际 N、G 与所选活动，给三种参与路径，不设资格门槛。','[ ] 按器材公式清点盘、棋子、笔、所选题卡和空表；奇数成员安排轮换。','[ ] 100% 试印校准/套准页，检查纸盘与棋子直径、可见坐标、灰度标记。','[ ] 微课/卡背先隐藏答案；核对题号、r1、先手、目标和真实完整盘边。','[ ] 若用程序：先看已可离线，再断网刷新演练；另存 PDF/PPTX。','[ ] 准备实体替代；备份当前本机记录，确认更新与清除范围。']),H('活动中'),L(['[ ] 接待新成员并允许观战/直接对弈；熟练成员只给小提示，不代下。','[ ] 先预测再验证；开放/未知分支保留讨论，争议保留原盘和历史。']),H('收尾'),L(['[ ] 未完局保存或写棋谱，不强造赢家；死子确认与预览区分。','[ ] 自愿三项反馈可跳过；下一人拿空表，不展示前人原文。','[ ] 记录实际问题与资源版本，未观察项保留空白；不编造实施成效。','[ ] 查看原文后才考虑导出；不默认发送，存好纸稿与本机备份。','[ ] 清点棋子、纸环、笔与题卡，擦除覆片，归还棋具。']),P('实际缺件或故障 ______________________________________________\n下次准备事项 __________________________________________________'))

feedback=page('成员自愿反馈空表',P('约半分钟，全部可不答，也可直接交空表或跳过。请避免填写自己或他人的姓名、班级、联系方式等个人资料。'),P('活动编号 __________________  实际活动日期 __________________'),H('今天活动难度感受'),P('[ ] 偏简单    [ ] 刚合适    [ ] 偏难    [ ] 不回答'),H('最喜欢的活动'),P('[ ] 找气与数气    [ ] 提子实验室    [ ] 双人吃子\n[ ] 完整对弈      [ ] 自选题        [ ] 同伴合作\n[ ] 其他          [ ] 不回答'),H('希望下次改变什么'),P('________________________________________________________________\n\n________________________________________________________________\n\n________________________________________________________________'),P('[ ] 愿意把这份匿名答复用于教师汇总导出（默认不选）。'),P('可选：相关题目 ID/revision __________\n问题类型 [ ] 题目过难 [ ] 说明不清 [ ] 程序错误 [ ] 其他'),P('一人一张、分开回收。共享设备/纸表的答复数不是独立人数证明；主持人不得代填导出同意。'))

observation=page('教师观察空表',P('仅记录真实看见或听见的事。没有观察就留空或写未观察；不做棋力、智力或心理诊断。'),P('活动编号 __________  日期 __________  资源 ID/revision __________'),T(['观察点','实际事实或原话 可留空'],[[t,'\n\n'] for t in ['提出理由','考虑对手回应','修改原想法','给同伴解释机会','规则或器材问题']],[38,136]),P('观察位置/情境 _________________________________________________\n\n组织者做了什么 _________________________________________________\n\n仍待核查的问题 _________________________________________________\n\n下一次可试的调整 _______________________________________________'),P('这些记录只反映所观察的片段，不自动代表整场合作质量或学习成效。'))

revision=page('资源修订与过程记录空表',P('活动编号 __________  记录编号 __________  日期 __________'),P('观察到的问题 __________________________________________________\n证据引用与实际原文 ______________________________________________\n拟改资源 ID/revision __________  修改后草稿 ID/revision __________\n修改内容 ________________________________________________________\n修改日期 __________  验证方法 __________________________________\n实际实施日期 __________  实施后观察 ______________________________\n下一步 __________________________________________________________'),P('状态自行选择 [ ] 计划 [ ] 已修订 [ ] 已实际实施\n未实施或未收反馈时，实施后观察保持未观察，不自动宣布改进有效。'),H('贡献与反思'),P('本人设计/修改 __________________________________________________\n同伴实际贡献 ____________________________________________________\nAI 协助 _________________________________________________________\n参考资料 ________________________________________________________\n真实实施内容 ____________________________________________________\n本人实际投入分钟 __________（不自动换算 CAS 工时）'),P('反思提问：实际遇到什么？证据是什么？改变了什么？哪些仍待验证？\n________________________________________________________________'),P('保留修改前后草稿；正式冻结题不在本活动中覆盖。若将来改正式资源，提升 revision、重新验证并同步相关纸卡/课件。'))

newcomer=page('新成员快速接待卡',P('欢迎随时加入。可以找气、看一手、和伙伴下棋或先观战，没有通关门槛。'),H('主持人可以这样说'),L(['“你想先看一手，还是找个伙伴直接玩？”','“棋子放在交叉点，沿线上下左右旁边的空点叫气。”','“你先选点，我可以在你需要时给一个小提示。”','“不想回答或填反馈也没关系。”']),H('给成员的三个选择'),P('找气：LIB-001 r1，正/背第 1 页，答案册第 4 页；先猜再看。\n小游戏：9 路黑先，先提到棋子即停止；它是入门小游戏。\n直接对弈：完整 9 路默认贴目 0，轮流落子，也可以停一手。'),H('最少需要'),P('一个完整棋盘、黑白棋子、愿意同行的伙伴；规则卡放桌边。没有屏幕一样能玩，主持人可用指尖/纸环说明气。'),P('给已有对局的新加入者先问双方是否愿意让其观察或接替；否则另开一盘。觉得太难可换找气，觉得已掌握可选后续题、13 路或自由对弈。'),P('两停进入确认，不自动删死子；有争议保留棋并实际继续。先提小游戏的胜利不等于完整围棋胜利。'))

history=page('实体对局逐手记录空表',P('对局编号 __________ 日期 __________ 尺寸 __________ 贴目 __________\n首手执色 __________ 起点 [ ] 空盘 [ ] 另附完整布局与先手\n摆棋转对局须核查各块有气；当前布局建立新历史，此前历史未知。'),P('A–T 列跳过 I，行从下往上；9 路左上 A9、右下 J1。每一合法动作记一行，place 记坐标；pass/resign 不写落点。无效动作不换手、不当停手。'),T(['手数','执色','落点或停手/认输','被提坐标','备注'],[['','','','',''] for _ in range(14)],[14,15,63,42,40]),P('起点图编号 __________  本页承接第 __________ 手\n两停确认：双方死子/中立共识另记，注释不移走原盘；争议恢复原盘，保留本记录，连续停手数归零。'),P('重复核对：依起点与记录重摆，比较提子后的完整布局与包括起点在内的先前布局，只比较棋子；停手例外。无法核实时先暂停并保留记录，不猜合法性或赢家。续页可复印空表，手数接续，回退时原分支另标作废。'))

manual=[intro,prep,rules,digital]+[session_page(s) for s in D['sessions']]+[
 page('四种可重复玩法',P('这些是社团教学玩法与任务目标，不是官方双人围棋赛制。成员可自由选择，不用速度或胜负作为唯一奖励。'),*[b for row in D['games'] for b in [H(row[0]),P(row[1]+'。'+row[2]),P(row[3])]],P('任务站可设 S＝min(4,G)，也可同桌换题。站点采用原题尺寸/先手；只有 9 路器材时选择完整 9 路题或明确标示的完整 5 路卡题，不用裁切偷换盘边。')),
 page('常见情况怎样处理',T(['遇到的情况','可以立即采取的回应'],D['dilemmas'],[31,143]),P('主持人可以说“先把盘面留住，我们一起看一手”。必要时暂停这局，允许成员转别的活动；保留问题比编造确定结论更有用。')),page('程序与完全无屏替代表',P('两列执行的是同一目标。无个人设备、投影不可用或首次未缓存时，直接用右列完成核心步骤。'),T(['活动','程序版','无屏版'],D['substitutions'],[28,52,94])),peer_page,records,
 page('纸面棋具和规则索引',paper['blocks'][0],T(['短码','文件与用途','页码'],print_rows[:10],[14,143,17])),
 page('题卡答案和演示盘索引',T(['短码','文件与用途','页码'],print_rows[10:],[14,143,17]),paper['blocks'][-1]),
 page('微课文件与页码索引',digital_index['blocks'][0],digital_index['blocks'][1]),
 page('已安装程序入口索引',digital_index['blocks'][2],digital_index['blocks'][3]),source_page]
supports=[('01_screen_alternatives','T23-ALTERNATIVES',manual[12]),('02_peer_teaching','T23-PEER-CARD',peer_page),('03_prepare_close','T23-CHECKLIST',checklist),('04_voluntary_feedback','T23-FEEDBACK-TEMPLATE',feedback),('05_teacher_observation','T23-OBSERVATION-TEMPLATE',observation),('06_revision_record','T23-REVISION-TEMPLATE',revision),('07_new_member','T23-NEW-MEMBER-CARD',newcomer),('08_paper_game_log','T23-GAME-LOG-TEMPLATE',history)]
docs=[dict(name='T23_facilitator_manual',**ident('T23-MANUAL'),pages=manual)]
docs += [dict(name=f'T23_activity_{i:02}',**ident(s['id']),pages=[session_page(s)]) for i,s in enumerate(D['sessions'],1)]
docs += [dict(name='T23_'+name,**ident(id),pages=[p]) for name,id,p in supports]

def font(run,size=11):
 run.font.name='Noto Sans SC';run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
 run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans SC')
def paragraph(doc,text,style=None,size=11):
 p=doc.add_paragraph(style=style);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=Pt(14)
 for line_idx,line in enumerate(text.split('\n')):
  if line_idx:p.add_run().add_break()
  font(p.add_run(line),size)
 return p
def table(doc,b):
 t=doc.add_table(rows=1,cols=len(b['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=b.get('widths') or [174/len(b['headers'])]*len(b['headers'])
 for col,w in zip(t.columns,widths):col.width=Mm(w)
 for n,v in enumerate(b['headers']):t.rows[0].cells[n].text=v
 for row in b['rows']:
  cells=t.add_row().cells
  for n,v in enumerate(row):cells[n].text=str(v)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for n,c in enumerate(row.cells):
   c.width=Mm(widths[n]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   props=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
   props.append(mar);borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   props.append(borders)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E4EDF1');props.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=Pt(14)
    for r in p.runs:font(r,10.5);r.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

def write_metadata(path,d):
 ns='http://schemas.openxmlformats.org/officeDocument/2006/custom-properties';vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
 custom=ET.Element('{'+ns+'}Properties',nsmap={None:ns,'vt':vt})
 for i,(key,value) in enumerate({k:d[k] for k in ['schemaVersion','id','revision','rulesetId']}.items(),2):
  el=ET.SubElement(custom,'{'+ns+'}property',fmtid='{D5CDD505-2E9C-101B-9397-08002B2CF9AE}',pid=str(i),name=key);ET.SubElement(el,'{'+vt+'}lpwstr').text=str(value)
 with zipfile.ZipFile(path) as z:members={n:z.read(n) for n in z.namelist()}
 ct=ET.fromstring(members['[Content_Types].xml']);ET.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/docProps/custom.xml',ContentType='application/vnd.openxmlformats-officedocument.custom-properties+xml')
 rel=ET.fromstring(members['_rels/.rels']);ET.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdT23Metadata',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties',Target='docProps/custom.xml')
 members['[Content_Types].xml']=ET.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True);members['_rels/.rels']=ET.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True);members['docProps/custom.xml']=ET.tostring(custom,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,v in members.items():z.writestr(n,v)

def markdown(d):
 lines=['---']+[f'{k}: {d[k]}' for k in ['schemaVersion','id','revision','rulesetId']]+['resultVersion: '+D['resultVersion'],'---','']
 for p in d['pages']:
  lines+=['# '+p['title'],'']
  for b in p['blocks']:
   if b['type']=='paragraph':lines+=[b['text'].replace('\n','  \n'),'']
   elif b['type']=='heading':lines+=['## '+b['text'],'']
   elif b['type']=='list':lines+=['- '+x for x in b['items']]+['']
   else:
    lines+=['| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
    lines+=['| '+' | '.join(str(x).replace('\n','<br>') for x in row)+' |' for row in b['rows']]+['']
 if d['id']=='T23-MANUAL':
  lines+=['## 原材料文件链接','']+[f'- {a} [{label}]({os.path.relpath(R/path,OUT)}) 第 {pages} 页' for a,label,path,pages in resources]
 return '\n'.join(lines)+'\n'

for d in docs:
 doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
 sec.left_margin=sec.right_margin=Mm(18);sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.footer_distance=Mm(7)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=doc.styles[name];st.font.name='Noto Sans SC';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans SC')
 doc.styles['Normal'].font.size=Pt(11)
 for element in doc.styles.element.xpath('.//w:pBdr'):
  element.getparent().remove(element)
 footer=sec.footer.paragraphs[0];font(footer.add_run('围棋社活动资源  '+d['id']+' r1  |  '),8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 for i,p in enumerate(d['pages']):
  # Explicit title page breaks avoid trailing empty break paragraphs.
  title=paragraph(doc,p['title'],'Title' if i==0 else 'Heading 1',21 if i==0 else 18);title.paragraph_format.space_after=Pt(8);title.paragraph_format.page_break_before=i>0
  for b in p['blocks']:
   if b['type']=='paragraph':paragraph(doc,b['text'])
   elif b['type']=='heading':paragraph(doc,b['text'],'Heading 2',13)
   elif b['type']=='list':
    for x in b['items']:paragraph(doc,'• '+x)
   else:table(doc,b)
 doc.core_properties.title=d['pages'][0]['title'];doc.core_properties.author='围棋社资源制作';doc.core_properties.subject='可调整活动设计与空白记录模板';doc.core_properties.keywords=d['id']+' '+D['resultVersion']+' GO-CLUB-AREA-1'
 target=OUT/(d['name']+'.docx');doc.save(target);write_metadata(target,d)
 (OUT/(d['name']+'.md')).write_text(markdown(d),encoding='utf-8')

models=dict(**ident('T23-PAGE-MODEL'),kind='MaterialPageModel',extensionScope='materials-only',documents=docs)
(K/'source/page-model.json').write_text(json.dumps(models,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
refs=dict(**ident('T23-RESOURCE-MAP'),baseVersion=D['baseVersion'],resultVersion=D['resultVersion'],paperResources=[dict(code=a,label=b,path=c,pages=e) for a,b,c,e in resources],slides=[dict(number=n,pptx=f'materials/micro-lessons/output/T22_MicroLesson_0{n}_r1.pptx',pdf=f'materials/micro-lessons/output/T22_MicroLesson_0{n}_r1.pdf',pages=5) for n in range(1,6)],modules=[dict(label=a,id=b,route=c) for a,b,c in module_names],puzzles=sorted({x for s in D['sessions'] for x in s['puzzles']}),fixtures=['END-002','END-006'])
(K/'source/resource-map.json').write_text(json.dumps(refs,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
templates=dict(**ident('T23-EMPTY-FORM-MODEL'),kind='MaterialTemplates',notSubmittedRecords=True,forms=[dict(id=id,revision=1,name=name,values={},submitted=False) for name,id,p in supports if 'TEMPLATE' in id])
(K/'source/empty-forms.json').write_text(json.dumps(templates,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Built {len(docs)} DOCX and Markdown documents; manual {len(manual)} intended pages; support and activity sheets one page each.')
